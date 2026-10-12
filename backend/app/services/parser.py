"""PDF/DOCX parsing that survives messy real documents. Never raises on odd content;
raises ParseError (with a user-friendly message) only for unusable files."""
import io
import re
import statistics
from dataclasses import dataclass, field

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 80
HEADING_RE = re.compile(r"^(\d{1,2}(\.\d{1,2}){0,3}\.?)\s+[A-Z][^\n]{2,90}$")
SCANNED_MIN_CHARS = 40


class ParseError(Exception):
    pass


@dataclass
class Block:
    text: str
    heading: bool = False


@dataclass
class ParsedPage:
    page: int | None
    blocks: list[Block] = field(default_factory=list)
    ocr: bool = False

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


@dataclass
class ParsedDoc:
    file_type: str
    pages: list[ParsedPage]
    page_count: int | None
    scanned_pages: list[int] = field(default_factory=list)  # 1-based pages needing OCR

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


def sniff(data: bytes) -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        return "docx"
    raise ParseError("Unsupported file. Please upload a PDF or DOCX.")


def _clean(t: str) -> str:
    t = t.replace("\x00", "").replace(" ", " ")
    return re.sub(r"[ \t]+", " ", t).strip()


def parse_pdf(data: bytes) -> ParsedDoc:
    import pymupdf
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:  # noqa: BLE001
        raise ParseError("This PDF is corrupt or unreadable.") from e
    if doc.needs_pass or doc.is_encrypted:
        raise ParseError("This PDF is password-protected. Remove the password and try again.")
    if doc.page_count == 0:
        raise ParseError("This PDF has no pages.")
    if doc.page_count > MAX_PAGES:
        raise ParseError(f"This PDF has {doc.page_count} pages; the limit is {MAX_PAGES}.")

    # body font size = median span size across the document
    sizes: list[float] = []
    raw_pages = []
    for page in doc:
        try:
            d = page.get_text("dict")
        except Exception:  # noqa: BLE001
            d = {"blocks": []}
        raw_pages.append(d)
        for b in d.get("blocks", []):
            for ln in b.get("lines", []):
                for sp in ln.get("spans", []):
                    if sp.get("text", "").strip():
                        sizes.append(round(sp["size"], 1))
    body = statistics.median(sizes) if sizes else 10.0

    pages: list[ParsedPage] = []
    scanned: list[int] = []
    for i, d in enumerate(raw_pages, start=1):
        pp = ParsedPage(page=i)
        for b in d.get("blocks", []):
            if b.get("type") != 0:
                continue
            for ln in b.get("lines", []):
                spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                text = _clean("".join(s["text"] for s in spans))
                if not text:
                    continue
                size = max(s["size"] for s in spans)
                bold = all((s.get("flags", 0) & 16) or "bold" in s.get("font", "").lower() for s in spans)
                is_head = len(text) <= 100 and not text.endswith((".", ",", ";")) and (
                    (HEADING_RE.match(text) is not None and len(text.split()) <= 10)
                    or (size >= body * 1.15 and len(text.split()) <= 12)
                    or (bold and len(text.split()) <= 8 and not text.endswith((".", ",", ";")))
                )
                pp.blocks.append(Block(text, is_head))
        if len(pp.text.strip()) < SCANNED_MIN_CHARS:
            scanned.append(i)
        pages.append(pp)
    return ParsedDoc("pdf", _merge_lines(pages), doc.page_count, scanned)


def _merge_lines(pages: list[ParsedPage]) -> list[ParsedPage]:
    """Join consecutive non-heading lines into paragraph-ish blocks (split on sentence-final lines)."""
    for p in pages:
        merged: list[Block] = []
        buf: list[str] = []

        def flush():
            if buf:
                merged.append(Block(" ".join(buf)))
                buf.clear()
        for b in p.blocks:
            if b.heading:
                flush()
                merged.append(b)
            else:
                buf.append(b.text)
                if b.text.endswith((".", ":", "?", "!")) and sum(len(x) for x in buf) > 200:
                    flush()
        flush()
        p.blocks = merged
    return pages


def parse_docx(data: bytes) -> ParsedDoc:
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ParseError("This DOCX is corrupt or unreadable.") from e
    pp = ParsedPage(page=None)
    for para in d.paragraphs:
        t = _clean(para.text)
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        pp.blocks.append(Block(t, style.startswith("heading") or style == "title" or bool(HEADING_RE.match(t))))
    for table in d.tables:
        for row in table.rows:
            cells = [_clean(c.text) for c in row.cells if _clean(c.text)]
            if cells:
                pp.blocks.append(Block(" | ".join(cells)))
    if not pp.text.strip():
        raise ParseError("No readable text was found in this DOCX.")
    return ParsedDoc("docx", [pp], None, [])


def parse(data: bytes) -> ParsedDoc:
    if len(data) > MAX_BYTES:
        raise ParseError("File is larger than 10 MB.")
    return parse_pdf(data) if sniff(data) == "pdf" else parse_docx(data)


def render_page_png(data: bytes, page_no: int, zoom: float = 2.0) -> bytes:
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    return doc[page_no - 1].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")


def find_highlights(data: bytes, page_no: int, chunk_text: str, max_phrases: int = 6) -> dict:
    """Rects for the cited passage on a page, from several short phrases of the chunk."""
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    page = doc[page_no - 1]
    words = chunk_text.split()
    phrases: list[str] = []
    n = 6
    for start in range(0, max(1, len(words) - n + 1), max(n, len(words) // max_phrases or 1)):
        ph = " ".join(words[start:start + n])
        if len(ph) >= 12:
            phrases.append(ph)
        if len(phrases) >= max_phrases:
            break
    rects: list[list[float]] = []
    for ph in phrases:
        for r in page.search_for(ph):
            rects.append([round(r.x0, 1), round(r.y0, 1), round(r.x1, 1), round(r.y1, 1)])
    return {"page": page_no, "page_width": page.rect.width, "page_height": page.rect.height, "rects": rects}
